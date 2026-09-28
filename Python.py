# -*- coding: utf-8 -*-
"""
Analyse du taux de chômage au Maroc (HCP, 2014-2025)
Question : pourquoi le taux de chômage augmente-t-il au Maroc ?

Le script :
  1. importe le fichier Excel brut du HCP,
  2. le transforme en tableau propre (data/taux_chomage_propre.csv),
  3. génère 3 graphiques dans le dossier graphiques/,
  4. affiche quelques indicateurs chiffrés utilisés dans le rapport.

Utilisation :  python analyse_chomage.py
Dépendances :  pip install openpyxl pandas matplotlib
"""

import os
from openpyxl import load_workbook
import pandas as pd
import matplotlib.pyplot as plt

FICHIER_BRUT = "data/taux_chomage_hcp_brut.xlsx"
FICHIER_PROPRE = "data/taux_chomage_propre.csv"
DOSSIER_GRAPHIQUES = "graphiques"


# ---------------------------------------------------------------
# ÉTAPE 1 : importer le fichier Excel et le mettre à plat
# ---------------------------------------------------------------
def importer_donnees(chemin):
    """Lit le fichier HCP et renvoie un DataFrame propre.

    Le fichier brut a une mise en forme "à cellules fusionnées" :
    le milieu (National/Urbain/Rural) et la tranche d'âge ne sont écrits
    qu'une seule fois puis laissés vides. On les reporte donc ligne
    après ligne avec une boucle.
    """
    feuille = load_workbook(chemin, data_only=True).active
    lignes = list(feuille.iter_rows(values_only=True))

    # 1re ligne : ["Périodes", None, None, 2025, 2024, ..., 2014]
    annees = []
    for valeur in lignes[0][3:]:
        annees.append(int(valeur))

    enregistrements = []
    milieu = None
    age = None

    # On saute les 2 lignes d'en-tête (Périodes / Milieu-Age-Sexe)
    for ligne in lignes[2:]:
        if ligne[0] is not None:   # nouvelle valeur de milieu
            milieu = ligne[0]
        if ligne[1] is not None:   # nouvelle tranche d'âge
            age = ligne[1]
        sexe = ligne[2]
        if sexe is None:
            continue

        # une ligne du fichier = une valeur par année
        for i in range(len(annees)):
            taux = ligne[3 + i]
            enregistrements.append({
                "annee": annees[i],
                "milieu": milieu,
                "age": age,
                "sexe": sexe,
                "taux_chomage": float(taux),
            })

    tableau = pd.DataFrame(enregistrements)
    tableau = tableau.sort_values(["milieu", "age", "sexe", "annee"])
    return tableau


def serie(tableau, milieu, age, sexe):
    """Extrait une série chronologique (années, taux) pour un cas donné."""
    sous_tableau = tableau[
        (tableau["milieu"] == milieu)
        & (tableau["age"] == age)
        & (tableau["sexe"] == sexe)
    ].sort_values("annee")
    return sous_tableau["annee"].tolist(), sous_tableau["taux_chomage"].tolist()


def habiller(ax, titre, ylabel="Taux de chômage (%)"):
    """Mise en forme commune à tous les graphiques."""
    ax.set_title(titre, fontsize=12, fontweight="bold")
    ax.set_xlabel("Année")
    ax.set_ylabel(ylabel)
    ax.grid(alpha=0.3)
    ax.legend()


# ---------------------------------------------------------------
# ÉTAPE 2 : les graphiques
# ---------------------------------------------------------------
def graphique_1(tableau):
    """Taux global : national, urbain, rural (15 ans et plus)."""
    fig, ax = plt.subplots(figsize=(9, 5))
    couleurs = {"National": "green", "Urbain": "blue", "Rural": "red"}

    for milieu in ["National", "Urbain", "Rural"]:
        x, y = serie(tableau, milieu, "15 ans et plus", "Total")
        ax.plot(x, y, marker="o", label=milieu, color=couleurs[milieu])
        # on écrit la dernière valeur au bout de la courbe
        ax.annotate(str(y[-1]), (x[-1], y[-1]), textcoords="offset points",
                    xytext=(6, 0), fontsize=9)

    ax.axvspan(2019.5, 2020.5, color="grey", alpha=0.15)
    ax.text(2020, 0.02, "Covid-19", ha="center", va="bottom", fontsize=8,
            transform=ax.get_xaxis_transform())
    habiller(ax, "Taux de chômage au Maroc (15 ans et plus), 2014-2025")
    fig.text(0.01, 0.01, "Source : HCP (bds.hcp.ma/data/1.2)", fontsize=8)
    fig.savefig(f"{DOSSIER_GRAPHIQUES}/graphique1_taux_global.png", dpi=150,
                bbox_inches="tight")
    plt.close(fig)


def graphique_2(tableau):
    """Chômage selon la tranche d'âge (niveau national, total)."""
    fig, ax = plt.subplots(figsize=(9, 5))
    tranches = ["15-24 ans", "25-34 ans", "35-44 ans", "45 ans et plus"]

    for tranche in tranches:
        x, y = serie(tableau, "National", tranche, "Total")
        ax.plot(x, y, marker="o", label=tranche)

    habiller(ax, "Le chômage touche surtout les jeunes (national, 2014-2025)")
    fig.text(0.01, 0.01, "Source : HCP (bds.hcp.ma/data/1.2)", fontsize=8)
    fig.savefig(f"{DOSSIER_GRAPHIQUES}/graphique2_par_age.png", dpi=150,
                bbox_inches="tight")
    plt.close(fig)


def graphique_3(tableau):
    """Écart hommes / femmes (national, 15 ans et plus)."""
    fig, ax = plt.subplots(figsize=(9, 5))

    for sexe, couleur in [("Masculin", "steelblue"), ("Feminin", "crimson")]:
        x, y = serie(tableau, "National", "15 ans et plus", sexe)
        ax.plot(x, y, marker="o", label=sexe.replace("Feminin", "Féminin"),
                color=couleur)

    habiller(ax, "Chômage des hommes et des femmes (15 ans et plus), 2014-2025")
    fig.text(0.01, 0.01, "Source : HCP (bds.hcp.ma/data/1.2)", fontsize=8)
    fig.savefig(f"{DOSSIER_GRAPHIQUES}/graphique3_hommes_femmes.png", dpi=150,
                bbox_inches="tight")
    plt.close(fig)


# ---------------------------------------------------------------
# ÉTAPE 3 : indicateurs chiffrés pour le rapport
# ---------------------------------------------------------------
def afficher_indicateurs(tableau):
    print("\n=== Évolution du taux de chômage national (15 ans et plus) ===")
    x, y = serie(tableau, "National", "15 ans et plus", "Total")
    for i in range(len(x)):
        if i == 0:
            print(f"{x[i]} : {y[i]} %")
        else:
            variation = round(y[i] - y[i - 1], 1)
            print(f"{x[i]} : {y[i]} %  ({variation:+} pt vs {x[i-1]})")

    print("\n=== Variations 2019 -> 2025 (points de %) ===")
    cas = [
        ("National", "15 ans et plus", "Total"),
        ("Urbain", "15 ans et plus", "Total"),
        ("Rural", "15 ans et plus", "Total"),
        ("National", "15-24 ans", "Total"),
        ("National", "15 ans et plus", "Feminin"),
        ("National", "15 ans et plus", "Masculin"),
    ]
    for milieu, age, sexe in cas:
        x, y = serie(tableau, milieu, age, sexe)
        dico = dict(zip(x, y))
        ecart = round(dico[2025] - dico[2019], 1)
        print(f"{milieu:9s} {age:15s} {sexe:9s}: {dico[2019]} -> {dico[2025]}  ({ecart:+} pt)")


# ---------------------------------------------------------------
# PROGRAMME PRINCIPAL
# ---------------------------------------------------------------
if __name__ == "__main__":
    os.makedirs(DOSSIER_GRAPHIQUES, exist_ok=True)

    donnees = importer_donnees(FICHIER_BRUT)
    donnees.to_csv(FICHIER_PROPRE, index=False, encoding="utf-8")
    print(f"{len(donnees)} lignes importées -> {FICHIER_PROPRE}")

    graphique_1(donnees)
    graphique_2(donnees)
    graphique_3(donnees)
    print(f"3 graphiques enregistrés dans {DOSSIER_GRAPHIQUES}/")

    afficher_indicateurs(donnees)
